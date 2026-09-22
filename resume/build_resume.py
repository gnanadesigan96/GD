"""Build ATS-friendly PDF and DOCX versions of the resume.

Usage:  python3 build_resume.py
Needs:  pip install reportlab python-docx

ATS rules followed in both outputs:
  * single column, no tables, text boxes, images, icons or headers/footers
  * standard section headings (Professional Summary, Skills, ...)
  * real text bullets and plain-text contact details
  * standard fonts (Helvetica in the PDF, Arial in the DOCX)
"""

import html
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Flowable, HRFlowable, Paragraph, SimpleDocTemplate

import resume_data as R

HERE = os.path.dirname(os.path.abspath(__file__))
BASENAME = "Athi_Shree_V_Resume"
ACCENT = "1F4E79"


def e(text):
    return html.escape(text)


# --------------------------------------------------------------------------- PDF
# Built with ReportLab using the standard Helvetica fonts (simple Type1 fonts,
# WinAnsi encoding, classic xref table). This is the most widely parsable PDF
# structure; browser-printed PDFs use embedded CID fonts that some ATS reject.
class DatedLine(Flowable):
    """One line with bold text on the left and a right-aligned date."""

    def __init__(self, left, right, size, color, space_before=0, rest=""):
        super().__init__()
        self.left, self.right, self.size, self.color = left, right, size, color
        self.rest = rest  # optional regular-weight text after the bold part
        self.spaceBefore = space_before

    def wrap(self, avail_w, avail_h):
        self.width = avail_w
        return avail_w, self.size * 1.3

    def draw(self):
        c = self.canv
        c.setFont("Helvetica-Bold", self.size)
        c.setFillColor(self.color)
        c.drawString(0, 2, self.left)
        if self.rest:
            c.setFont("Helvetica", self.size)
            c.drawString(c.stringWidth(self.left, "Helvetica-Bold", self.size), 2, self.rest)
            c.setFont("Helvetica-Bold", self.size)
        c.setFillColor(colors.HexColor("#333333"))
        c.drawRightString(self.width, 2, self.right)


def build_pdf(pdf_path):
    accent = colors.HexColor("#" + ACCENT)
    ink = colors.HexColor("#1a1a1a")
    base = dict(fontName="Helvetica", fontSize=9.2, leading=11.2, textColor=ink)
    st = {
        "name": ParagraphStyle("name", fontName="Helvetica-Bold", fontSize=21, leading=24,
                               alignment=TA_CENTER, textColor=accent),
        "title": ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=10.4, leading=13,
                                alignment=TA_CENTER, textColor=colors.HexColor("#333333")),
        "contact": ParagraphStyle("contact", **{**base, "fontSize": 9.3}, alignment=TA_CENTER,
                                  spaceAfter=2),
        "h2": ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=10.8, leading=13,
                             textColor=accent, spaceBefore=5),
        "body": ParagraphStyle("body", **base, alignment=TA_JUSTIFY),
        "skill": ParagraphStyle("skill", **base, spaceAfter=1),
        "bullet": ParagraphStyle("bullet", **base, leftIndent=11, bulletIndent=2, spaceAfter=0.8),
    }

    def heading(text):
        return [Paragraph(text.upper(), st["h2"]),
                HRFlowable(width="100%", thickness=1.2, color=accent, spaceBefore=1, spaceAfter=3)]

    def bullets(items):
        return [Paragraph(e(t), st["bullet"], bulletText="\u2022") for t in items]

    story = [
        Paragraph(e(R.NAME), st["name"]),
        Paragraph(e(R.TITLE), st["title"]),
        Paragraph(" | ".join(e(c) for c in R.CONTACT), st["contact"]),
        *heading("Professional Summary"),
        Paragraph(e(R.SUMMARY), st["body"]),
        *heading("Skills"),
    ]
    for k, v in R.SKILLS:
        story.append(Paragraph(f'<font name="Helvetica-Bold" color="#{ACCENT}">{e(k)}:</font> {e(v)}',
                               st["skill"]))
    story += heading("Professional Experience")
    story.append(DatedLine(f"{R.EMPLOYER_ROLE} | {R.EMPLOYER}", R.EMPLOYER_DATES, 10.2, ink))
    for p in R.PROJECTS:
        story.append(DatedLine(f"Client: {p['client']} | {p['role']}", p["dates"], 9.4, accent,
                               space_before=3))
        story += bullets(p["bullets"])
    story += heading("Certifications") + bullets(R.CERTIFICATIONS)
    story += heading("Education")
    for d, s_, y, g in R.EDUCATION:
        story.append(DatedLine(d, y, 9.3, ink, space_before=1, rest=f" | {s_} | {g}"))
    story += heading("Achievements") + bullets(R.ACHIEVEMENTS)

    doc = SimpleDocTemplate(pdf_path, pagesize=A4, leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=8 * mm, bottomMargin=6 * mm,
                            title=f"{R.NAME.title()} - Resume", author=R.NAME.title(),
                            subject="Resume", creator="build_resume.py")
    doc.build(story)


# --------------------------------------------------------------------------- DOCX
def _bottom_border(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "10"), ("space", "1"), ("color", ACCENT)):
        b.set(qn(f"w:{k}"), v)
    bdr.append(b)
    pPr.append(bdr)


def build_docx(path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.left_margin = sec.right_margin = Inches(0.6)
    sec.top_margin = sec.bottom_margin = Inches(0.5)
    accent = RGBColor.from_string(ACCENT)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.08

    def para(text="", bold=False, size=None, color=None, align=None, italic=False, after=0, before=0):
        p = doc.add_paragraph()
        if text:
            r = p.add_run(text)
            r.bold, r.italic = bold, italic
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.space_before = Pt(before)
        return p

    def heading(text):
        p = para(text.upper(), bold=True, size=11, color=accent, before=8, after=3)
        _bottom_border(p)

    def bullet(text):
        p = doc.add_paragraph(text, style="List Bullet")
        p.paragraph_format.space_after = Pt(1)

    def dated(left, right, color=None, before=2):
        # Right-aligned date via a tab stop (no tables, stays ATS-safe).
        p = para(before=before)
        p.paragraph_format.tab_stops.add_tab_stop(sec.page_width - sec.left_margin - sec.right_margin,
                                                  alignment=2)
        r = p.add_run(left)
        r.bold = True
        if color:
            r.font.color.rgb = color
        r2 = p.add_run("\t" + right)
        r2.bold = True
        return p

    para(R.NAME, bold=True, size=22, color=accent, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(R.TITLE, bold=True, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, before=2)
    para(" | ".join(R.CONTACT), size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)

    heading("Professional Summary")
    para(R.SUMMARY).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    heading("Skills")
    for k, v in R.SKILLS:
        p = para(after=1)
        r = p.add_run(k + ": ")
        r.bold = True
        r.font.color.rgb = accent
        p.add_run(v)

    heading("Professional Experience")
    dated(f"{R.EMPLOYER_ROLE} | {R.EMPLOYER}", R.EMPLOYER_DATES)
    for p in R.PROJECTS:
        dated(f"Client: {p['client']} | {p['role']}", p["dates"], color=accent, before=5)
        for b in p["bullets"]:
            bullet(b)

    heading("Certifications")
    for c in R.CERTIFICATIONS:
        bullet(c)

    heading("Education")
    for d, s, y, g in R.EDUCATION:
        p = dated(d, y)
        p.runs[0]._r.addnext(p.add_run(f" | {s} | {g}")._r)

    heading("Achievements")
    for a in R.ACHIEVEMENTS:
        bullet(a)

    doc.core_properties.title = f"{R.NAME.title()} - Resume"
    doc.core_properties.author = R.NAME.title()
    doc.save(path)


if __name__ == "__main__":
    build_pdf(os.path.join(HERE, BASENAME + ".pdf"))
    build_docx(os.path.join(HERE, BASENAME + ".docx"))
    print("Built", BASENAME + ".pdf / .docx")
